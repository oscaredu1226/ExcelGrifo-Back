from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
import re
import unicodedata

import openpyxl
import pandas as pd

from app.utils.money import normalize_money


TIPO_IZIPAY = "IZIPAY"
TIPO_FACTURACION = "FACTURACION"


FIELD_ALIASES: dict[str, dict[str, list[str]]] = {
    TIPO_IZIPAY: {
        "fecha_transaccion": ["fecha de transaccion", "fecha transaccion", "fec transaccion"],
        "hora_transaccion": ["hora de transaccion", "hora transaccion", "hora"],
        "importe": ["importe", "monto", "importe venta"],
        "importe_neto": ["importe neto", "neto", "monto neto"],
        "voucher": ["voucher", "nro voucher", "numero voucher"],
        "codigo_autorizacion": ["codigo de autorizacion", "cod autorizacion", "autorizacion"],
        "terminal": ["terminal", "pos", "codigo terminal"],
        "tipo_movimiento": ["tipo de movimiento", "movimiento"],
        "transaccion": ["transaccion", "tipo transaccion"],
        "estado": ["estado", "estado operacion"],
    },
    TIPO_FACTURACION: {
        "tipo_documento": ["tipo de comprobante", "tipo comprobante", "tipo documento", "tipo doc"],
        "serie": ["serie", "serie comprobante"],
        "numero": ["numero", "numero comprobante", "nro comprobante", "correlativo"],
        "fecha_emision": ["fecha de emision", "fecha emision", "fec emision"],
        "fecha_pago": ["fecha de pago", "fecha pago", "fec pago"],
        "documento_cliente": ["documento cliente", "doc cliente", "ruc", "dni", "numero documento"],
        "nombre_cliente": ["nombre cliente", "cliente", "razon social"],
        "importe": [
            "importe total del comprobante de pago",
            "importe total",
            "total comprobante",
            "monto total",
            "importe",
            "total",
        ],
    },
}


@dataclass(frozen=True)
class ExcelDetection:
    tipo: str
    sheet_name: str
    header_row: int
    columns: dict[str, str]


def validate_xlsx(filename: str) -> None:
    if not filename.lower().endswith(".xlsx"):
        raise ValueError("Solo se aceptan archivos .xlsx.")


def normalize_header(value: object) -> str:
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = text.lower().strip()
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def _matches(header: str, alias: str) -> bool:
    header_norm = normalize_header(header)
    alias_norm = normalize_header(alias)
    return header_norm == alias_norm or alias_norm in header_norm or header_norm in alias_norm


def _map_columns(headers: list[object], tipo: str) -> dict[str, str]:
    result: dict[str, str] = {}
    header_texts = [str(header).strip() for header in headers if header not in (None, "")]
    for field, aliases in FIELD_ALIASES[tipo].items():
        exact_match = next(
            (
                header
                for header in header_texts
                if any(normalize_header(header) == normalize_header(alias) for alias in aliases)
            ),
            None,
        )
        if exact_match:
            result[field] = exact_match
            continue

        for header in header_texts:
            if any(_matches(header, alias) for alias in aliases):
                result[field] = header
                break
    return result


def _score_columns(columns: dict[str, str], tipo: str) -> int:
    required = {
        TIPO_IZIPAY: {"fecha_transaccion", "importe", "voucher", "terminal", "codigo_autorizacion"},
        TIPO_FACTURACION: {"tipo_documento", "serie", "numero", "fecha_emision", "fecha_pago", "importe"},
    }
    return len(required[tipo].intersection(columns))


def detect_excel(path: Path) -> ExcelDetection | None:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        best: tuple[int, ExcelDetection] | None = None
        for row_number, row in enumerate(sheet.iter_rows(min_row=1, max_row=30, values_only=True), start=1):
            if not any(cell not in (None, "") for cell in row):
                continue

            izipay_columns = _map_columns(list(row), TIPO_IZIPAY)
            facturacion_columns = _map_columns(list(row), TIPO_FACTURACION)
            izipay_score = _score_columns(izipay_columns, TIPO_IZIPAY)
            facturacion_score = _score_columns(facturacion_columns, TIPO_FACTURACION)

            if izipay_score >= 3 and izipay_score > facturacion_score:
                candidate = ExcelDetection(TIPO_IZIPAY, sheet.title, row_number, izipay_columns)
                if best is None or izipay_score > best[0]:
                    best = (izipay_score, candidate)
            elif facturacion_score >= 4 and facturacion_score > izipay_score:
                candidate = ExcelDetection(TIPO_FACTURACION, sheet.title, row_number, facturacion_columns)
                if best is None or facturacion_score > best[0]:
                    best = (facturacion_score, candidate)
        return best[1] if best else None
    finally:
        workbook.close()


def identify_excel_file(path: Path) -> ExcelDetection:
    detection = detect_excel(path)
    if detection is None:
        raise ValueError("No se pudo identificar el archivo como Izipay o Facturacion.")
    return detection


def read_facturacion(path: Path) -> list[dict]:
    detection = identify_excel_file(path)
    if detection.tipo != TIPO_FACTURACION:
        raise ValueError("El archivo no corresponde a Facturacion.")
    df = pd.read_excel(
        path,
        sheet_name=detection.sheet_name,
        header=detection.header_row - 1,
        dtype=object,
        engine="openpyxl",
    )
    return _facturacion_rows(df, detection)


def read_izipay(path: Path) -> list[dict]:
    detection = identify_excel_file(path)
    if detection.tipo != TIPO_IZIPAY:
        raise ValueError("El archivo no corresponde a Izipay.")
    df = pd.read_excel(
        path,
        sheet_name=detection.sheet_name,
        header=detection.header_row - 1,
        dtype=object,
        engine="openpyxl",
    )
    return _izipay_rows(df, detection)


def _value(row: pd.Series, columns: dict[str, str], field: str) -> object | None:
    column = columns.get(field)
    if column is None or column not in row:
        return None
    value = row[column]
    if pd.isna(value):
        return None
    return value


def _clean_text(value: object | None) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and pd.isna(value):
        return None
    text = str(value).strip()
    if text.endswith(".0"):
        text = text[:-2]
    return text or None


def _clean_document_type(value: object | None) -> str | None:
    text = _clean_text(value)
    if text and text.isdigit() and len(text) == 1:
        return text.zfill(2)
    return text


def _parse_date(value: object | None) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    parsed = pd.to_datetime(value, errors="coerce", dayfirst=True)
    if pd.isna(parsed):
        return None
    return parsed.date()


def _parse_time(value: object | None) -> time | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.time().replace(microsecond=0)
    if isinstance(value, time):
        return value.replace(microsecond=0)
    parsed = pd.to_datetime(value, errors="coerce")
    if pd.isna(parsed):
        return None
    return parsed.time().replace(microsecond=0)


def _facturacion_rows(df: pd.DataFrame, detection: ExcelDetection) -> list[dict]:
    rows: list[dict] = []
    for index, row in df.iterrows():
        if row.isna().all():
            continue

        importe = normalize_money(_value(row, detection.columns, "importe"))
        if importe is None or importe <= 0:
            continue

        tipo_documento = _clean_document_type(_value(row, detection.columns, "tipo_documento"))
        if tipo_documento == "07":
            continue

        rows.append(
            {
                "fila_original": detection.header_row + 1 + int(index),
                "tipo_documento": tipo_documento,
                "serie": _clean_text(_value(row, detection.columns, "serie")),
                "numero": _clean_text(_value(row, detection.columns, "numero")),
                "fecha_emision": _parse_date(_value(row, detection.columns, "fecha_emision")),
                "fecha_pago": _parse_date(_value(row, detection.columns, "fecha_pago")),
                "documento_cliente": _clean_text(_value(row, detection.columns, "documento_cliente")),
                "nombre_cliente": _clean_text(_value(row, detection.columns, "nombre_cliente")),
                "importe": importe,
            }
        )
    return rows


def _is_valid_izipay(row_data: dict) -> bool:
    tipo_movimiento = normalize_header(row_data.get("tipo_movimiento"))
    transaccion = normalize_header(row_data.get("transaccion"))
    estado = normalize_header(row_data.get("estado"))

    if tipo_movimiento and "abono" not in tipo_movimiento:
        return False
    if transaccion and "gasolina" not in transaccion:
        return False
    if estado and "procesado" not in estado:
        return False
    return True


def _izipay_rows(df: pd.DataFrame, detection: ExcelDetection) -> list[dict]:
    rows: list[dict] = []
    for index, row in df.iterrows():
        if row.isna().all():
            continue

        importe = normalize_money(_value(row, detection.columns, "importe"))
        if importe is None or importe <= 0:
            continue

        data = {
            "fila_original": detection.header_row + 1 + int(index),
            "fecha_transaccion": _parse_date(_value(row, detection.columns, "fecha_transaccion")),
            "hora_transaccion": _parse_time(_value(row, detection.columns, "hora_transaccion")),
            "importe": importe,
            "importe_neto": normalize_money(_value(row, detection.columns, "importe_neto")),
            "voucher": _clean_text(_value(row, detection.columns, "voucher")),
            "codigo_autorizacion": _clean_text(_value(row, detection.columns, "codigo_autorizacion")),
            "terminal": _clean_text(_value(row, detection.columns, "terminal")),
            "tipo_movimiento": _clean_text(_value(row, detection.columns, "tipo_movimiento")),
            "transaccion": _clean_text(_value(row, detection.columns, "transaccion")),
            "estado": _clean_text(_value(row, detection.columns, "estado")),
        }
        if _is_valid_izipay(data):
            rows.append(data)
    return rows
