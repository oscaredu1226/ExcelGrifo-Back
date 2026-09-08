from copy import copy
from decimal import Decimal
from pathlib import Path
import re
from zipfile import ZIP_DEFLATED, ZipFile

from fastapi import HTTPException
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.orm import Session, joinedload

from app.models.comprobante import Comprobante
from app.models.emparejamiento import Emparejamiento
from app.models.operacion_izipay import OperacionIzipay
from app.services.conciliacion_service import get_conciliacion_or_404
from app.services.excel_service import TIPO_FACTURACION, TIPO_IZIPAY, identify_excel_file


EXPORTS_DIR = Path("exports")
UPLOADS_DIR = Path("uploads")

FACTURACION_EXTRA_COLUMNS = [
    {"header": "Medio de pago", "kind": "text_center", "width": 14},
    {"header": "Estado de match", "kind": "status", "width": 20},
    {"header": "Tipo de match", "kind": "text_center", "width": 16},
    {"header": "Fila Izipay", "kind": "integer", "width": 12},
    {"header": "Fecha transacción Izipay", "kind": "date", "width": 22},
    {"header": "Hora transacción Izipay", "kind": "time", "width": 20},
    {"header": "Importe Izipay", "kind": "money", "width": 15},
    {"header": "Voucher", "kind": "text", "width": 14},
    {"header": "Código de autorización", "kind": "text", "width": 22},
    {"header": "Terminal", "kind": "text", "width": 14},
]

IZIPAY_EXTRA_COLUMNS = [
    {"header": "Estado de match", "kind": "status", "width": 20},
    {"header": "Tipo de match", "kind": "text_center", "width": 16},
    {"header": "Fila Facturación", "kind": "integer", "width": 16},
    {"header": "Tipo comprobante", "kind": "text_center", "width": 17},
    {"header": "Serie", "kind": "text", "width": 12},
    {"header": "Número", "kind": "text", "width": 14},
    {"header": "Fecha emisión Facturación", "kind": "date", "width": 23},
    {"header": "Fecha pago Facturación", "kind": "date", "width": 22},
    {"header": "Importe Facturación", "kind": "money", "width": 18},
    {"header": "Documento cliente", "kind": "text", "width": 20},
    {"header": "Cliente", "kind": "text", "width": 32},
]

THIN_BORDER = Border(
    left=Side(style="thin", color="D0D5DD"),
    right=Side(style="thin", color="D0D5DD"),
    top=Side(style="thin", color="D0D5DD"),
    bottom=Side(style="thin", color="D0D5DD"),
)
HEADER_FILL = PatternFill("solid", fgColor="E5E7EB")
MATCH_FILL = PatternFill("solid", fgColor="DCFCE7")
WARNING_FILL = PatternFill("solid", fgColor="FEF3C7")
ERROR_FILL = PatternFill("solid", fgColor="FEE2E2")
NEUTRAL_FILL = PatternFill("solid", fgColor="F9FAFB")


def export_conciliacion(db: Session, conciliacion_id: int) -> Path:
    conciliacion = get_conciliacion_or_404(db, conciliacion_id)
    facturacion_path = UPLOADS_DIR / str(conciliacion_id) / "facturacion_original.xlsx"
    izipay_path = UPLOADS_DIR / str(conciliacion_id) / "izipay_original.xlsx"
    if not facturacion_path.exists() or not izipay_path.exists():
        raise HTTPException(status_code=400, detail="No existen los archivos originales de Facturacion e Izipay.")

    comprobantes = list(
        db.scalars(select(Comprobante).where(Comprobante.conciliacion_id == conciliacion_id)).all()
    )
    operaciones = list(
        db.scalars(select(OperacionIzipay).where(OperacionIzipay.conciliacion_id == conciliacion_id)).all()
    )
    emparejamientos = list(
        db.scalars(
            select(Emparejamiento)
            .where(Emparejamiento.conciliacion_id == conciliacion_id)
            .options(joinedload(Emparejamiento.comprobante), joinedload(Emparejamiento.operacion_izipay))
        ).all()
    )

    export_dir = EXPORTS_DIR / str(conciliacion_id)
    export_dir.mkdir(parents=True, exist_ok=True)

    facturacion_export = _export_facturacion_workbook(
        facturacion_path,
        conciliacion.archivo_facturacion,
        conciliacion.sucursal,
        conciliacion.mes,
        conciliacion.anio,
        comprobantes,
        {item.comprobante_id: item for item in emparejamientos},
        export_dir,
    )
    izipay_export = _export_izipay_workbook(
        izipay_path,
        conciliacion.archivo_izipay,
        conciliacion.sucursal,
        conciliacion.mes,
        conciliacion.anio,
        operaciones,
        {item.operacion_izipay_id: item for item in emparejamientos},
        export_dir,
    )

    zip_name = _export_filename(conciliacion.archivo_facturacion, conciliacion.sucursal, conciliacion.mes, conciliacion.anio, ".zip")
    zip_path = _unique_output_path(export_dir / zip_name)
    with ZipFile(zip_path, "w", compression=ZIP_DEFLATED) as zip_file:
        zip_file.write(facturacion_export, arcname=facturacion_export.name)
        zip_file.write(izipay_export, arcname=izipay_export.name)
    return zip_path


def _export_facturacion_workbook(
    facturacion_path: Path,
    original_name: str | None,
    sucursal: str,
    mes: int,
    anio: int,
    comprobantes: list[Comprobante],
    emparejamientos: dict[int, Emparejamiento],
    export_dir: Path,
) -> Path:
    detection = identify_excel_file(facturacion_path)
    if detection.tipo != TIPO_FACTURACION:
        raise HTTPException(status_code=400, detail="El archivo original de Facturacion no pudo identificarse.")

    workbook = load_workbook(facturacion_path)
    worksheet = workbook[detection.sheet_name]
    start_column = worksheet.max_column + 1
    _write_extra_headers(worksheet, detection.header_row, start_column, FACTURACION_EXTRA_COLUMNS)

    for comprobante in comprobantes:
        emparejamiento = emparejamientos.get(comprobante.id)
        if emparejamiento:
            operacion = emparejamiento.operacion_izipay
            values = [
                "IZIPAY",
                "MATCH",
                emparejamiento.tipo,
                operacion.fila_original,
                operacion.fecha_transaccion,
                operacion.hora_transaccion,
                operacion.importe,
                operacion.voucher,
                operacion.codigo_autorizacion,
                operacion.terminal,
            ]
        else:
            values = [None, comprobante.estado_conciliacion or "PENDIENTE", None, None, None, None, None, None, None, None]

        _write_extra_values(worksheet, comprobante.fila_original, start_column, FACTURACION_EXTRA_COLUMNS, values)

    _extend_auto_filter(worksheet, start_column + len(FACTURACION_EXTRA_COLUMNS) - 1)
    output_path = _unique_output_path(export_dir / _export_filename(original_name, sucursal, mes, anio))
    workbook.save(output_path)
    return output_path


def _export_izipay_workbook(
    izipay_path: Path,
    original_name: str | None,
    sucursal: str,
    mes: int,
    anio: int,
    operaciones: list[OperacionIzipay],
    emparejamientos: dict[int, Emparejamiento],
    export_dir: Path,
) -> Path:
    detection = identify_excel_file(izipay_path)
    if detection.tipo != TIPO_IZIPAY:
        raise HTTPException(status_code=400, detail="El archivo original de Izipay no pudo identificarse.")

    workbook = load_workbook(izipay_path)
    worksheet = workbook[detection.sheet_name]
    start_column = worksheet.max_column + 1
    _write_extra_headers(worksheet, detection.header_row, start_column, IZIPAY_EXTRA_COLUMNS)

    for operacion in operaciones:
        emparejamiento = emparejamientos.get(operacion.id)
        if emparejamiento:
            comprobante = emparejamiento.comprobante
            values = [
                "MATCH",
                emparejamiento.tipo,
                comprobante.fila_original,
                comprobante.tipo_documento,
                comprobante.serie,
                comprobante.numero,
                comprobante.fecha_emision,
                comprobante.fecha_pago,
                comprobante.importe,
                comprobante.documento_cliente,
                comprobante.nombre_cliente,
            ]
        elif operacion.estado_conciliacion == "BOLETA":
            values = ["BOLETA", "BOLETA", None, "BOLETA", None, None, None, None, operacion.importe, None, None]
        else:
            values = [operacion.estado_conciliacion or "PENDIENTE", None, None, None, None, None, None, None, None, None, None]

        _write_extra_values(worksheet, operacion.fila_original, start_column, IZIPAY_EXTRA_COLUMNS, values)

    _extend_auto_filter(worksheet, start_column + len(IZIPAY_EXTRA_COLUMNS) - 1)
    output_path = _unique_output_path(export_dir / _export_filename(original_name, sucursal, mes, anio))
    workbook.save(output_path)
    return output_path


def _write_extra_headers(worksheet: Worksheet, header_row: int, start_column: int, columns: list[dict]) -> None:
    for offset, column in enumerate(columns):
        cell = worksheet.cell(row=header_row, column=start_column + offset, value=column["header"])
        cell.font = Font(name="Aptos", size=10, bold=True, color="111827")
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        worksheet.column_dimensions[cell.column_letter].width = column["width"]

    worksheet.row_dimensions[header_row].height = max(worksheet.row_dimensions[header_row].height or 0, 30)


def _write_extra_values(
    worksheet: Worksheet,
    row: int,
    start_column: int,
    columns: list[dict],
    values: list[object | None],
) -> None:
    template_cell = worksheet.cell(row=row, column=max(1, start_column - 1))
    for offset, value in enumerate(values):
        column = columns[offset]
        cell = worksheet.cell(row=row, column=start_column + offset, value=_excel_value(value, column["kind"]))
        _copy_cell_style(template_cell, cell)
        _apply_extra_cell_format(cell, column["kind"])


def _copy_cell_style(source, target) -> None:
    if source.has_style:
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        target.number_format = source.number_format
        target.protection = copy(source.protection)
    target.border = copy(source.border) if source.has_style else THIN_BORDER


def _excel_value(value: object | None, kind: str) -> object | None:
    if value is None:
        return None
    if kind == "money" and isinstance(value, Decimal):
        return float(value)
    if kind in {"text", "text_center", "status"}:
        return str(value)
    return value


def _apply_extra_cell_format(cell, kind: str) -> None:
    cell.font = Font(name="Aptos", size=10, color="111827")
    cell.border = THIN_BORDER
    cell.alignment = Alignment(
        horizontal=_horizontal_alignment(kind),
        vertical="center",
        wrap_text=False,
    )

    if kind in {"text", "text_center", "status"}:
        cell.number_format = "@"
    elif kind == "integer":
        cell.number_format = "#,##0"
    elif kind == "date":
        cell.number_format = "dd/mm/yyyy"
    elif kind == "time":
        cell.number_format = "hh:mm:ss"
    elif kind == "money":
        cell.number_format = "#,##0.00"

    if kind == "status":
        _apply_status_fill(cell)


def _horizontal_alignment(kind: str) -> str:
    if kind in {"integer", "date", "time", "money"}:
        return "right"
    if kind in {"text_center", "status"}:
        return "center"
    return "left"


def _apply_status_fill(cell) -> None:
    value = str(cell.value or "").upper()
    if value in {"MATCH", "BOLETA"}:
        cell.fill = MATCH_FILL
    elif value in {"AMBIGUO", "SUGERENCIA", "PENDIENTE"}:
        cell.fill = WARNING_FILL
    elif value == "SIN_COINCIDENCIA":
        cell.fill = ERROR_FILL
    else:
        cell.fill = NEUTRAL_FILL


def _extend_auto_filter(worksheet: Worksheet, last_column: int) -> None:
    if not worksheet.auto_filter.ref:
        return

    min_col, min_row, _, max_row = range_boundaries(worksheet.auto_filter.ref)
    worksheet.auto_filter.ref = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(last_column)}{max_row}"


def _export_filename(original_name: str | None, sucursal: str, mes: int, anio: int, extension: str = ".xlsx") -> str:
    base_name = Path(original_name or f"Facturacion {sucursal} {mes} {anio}.xlsx").stem
    base_name = re.sub(r'[<>:"/\\|?*]+', "", base_name).strip(" .")
    return f"{base_name or 'Facturacion'}_MATCH{extension}"


def _unique_output_path(path: Path) -> Path:
    if not path.exists():
        return path

    for index in range(2, 1000):
        candidate = path.with_name(f"{path.stem}_{index}{path.suffix}")
        if not candidate.exists():
            return candidate

    raise HTTPException(status_code=409, detail="No se pudo generar un nombre disponible para la exportacion.")
